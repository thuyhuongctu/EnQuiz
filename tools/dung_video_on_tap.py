"""Dựng toàn bộ gói video ôn tập 5 chương (on-tap-video/ + assets/slides/)
từ hai nguồn: docs/kich_ban_on_tap.json (lời đọc) và docs/slide_noi_dung.json
(tiêu đề + gạch đầu dòng hiển thị trên từng slide — tóm tắt từ lời đọc, không
phải chép nguyên văn, theo đúng yêu cầu "không viết toàn văn trên slide" của
Điều 18 Quy định đào tạo trực tuyến và xây dựng học liệu số, QĐ
55/QĐ-ĐHSPKTVL-ĐT_ĐBCL ngày 07/7/2025).

Sinh ra năm thứ:
  - assets/slides/ch01..ch05/NNN.jpg (ảnh slide — dựng bằng HTML + Playwright,
    mẫu ở tools/_mau_slide_on_tap.py, KHÔNG qua LibreOffice/soffice vì soffice
    trong môi trường làm việc không chuyển được bất kỳ tệp nào sang PDF)
  - on-tap-video/kich-ban/KICH_BAN_THUYET_MINH__ON_TAP_EC1606.docx + .pdf
  - on-tap-video/slides/ON_TAP_5_CHUONG_EC1606.pdf (ghép toàn bộ ảnh slide
    thành một PDF — bản gốc để Huong dựng lại ảnh minh họa 3D đất sét; không
    phải tệp trình xem trong ứng dụng dùng, ảnh gốc trong assets/slides/ vẫn
    là nguồn cho ứng dụng)
  - on-tap-video/videos/CHUONG_N__ON_TAP.mp4 (năm video CÂM — chỉ ảnh slide,
    chưa có giọng đọc; xem on-tap-video/README.md)

Bản .pdf của kịch bản dựng thẳng bằng reportlab (không qua soffice), dùng
phông DejaVu Sans có sẵn trong hệ thống vì phông này đọc đủ dấu tiếng Việt.

Mỗi điểm đọc trong docs/kich_ban_on_tap.json ứng với ĐÚNG MỘT slide trong
docs/slide_noi_dung.json (ánh xạ 1-1 theo thứ tự — khác bản trước đó từng có
ngoại lệ ở Chương 1, bản hiện tại không còn ngoại lệ vì cả hai tệp được viết
lại cùng lúc với đúng số lượng khớp nhau).

Sửa lời đọc thì sửa docs/kich_ban_on_tap.json, sửa tiêu đề/gạch đầu dòng trên
slide thì sửa docs/slide_noi_dung.json, rồi chạy lại tập lệnh này — đừng sửa
tay ảnh slide hay các tệp kết quả trong on-tap-video/.

Chạy:  pip install python-docx reportlab   # nếu chưa có
       python3 tools/dung_video_on_tap.py
"""
import json
import os
import subprocess
import sys

import docx
from docx.shared import Pt
from PIL import Image
from playwright.sync_api import sync_playwright
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _mau_slide_on_tap as mau  # noqa: E402

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KICH_BAN_JSON = os.path.join(GOC, 'docs', 'kich_ban_on_tap.json')
SLIDE_NOI_DUNG_JSON = os.path.join(GOC, 'docs', 'slide_noi_dung.json')
THU_MUC_RA = os.path.join(GOC, 'on-tap-video')
THU_MUC_VIDEO = os.path.join(THU_MUC_RA, 'videos')
THU_MUC_KICH_BAN = os.path.join(THU_MUC_RA, 'kich-ban')
THU_MUC_SLIDES = os.path.join(THU_MUC_RA, 'slides')

PHONG_THUONG = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
PHONG_DAM = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
CHROME_PATH = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'

TOC_DO_DOC = 206.0  # từ/phút — đo thật ở ec6035tx-chuong-5 (ComDraft)
NGHI_CHUYEN_SLIDE = 1.0  # giây, theo quy ước ComDraft
THOI_LUONG_TOI_THIEU = 4.0

# (mã chương, khóa trong kich_ban_on_tap.json, số chương, tên chương đầy đủ)
CAC_CHUONG = [
    ('ch01', 'ON TAP CHUONG 1 - TONG QUAN VE KHOI SU DOANH NGHIEP', 1,
     'Tổng quan về khởi sự doanh nghiệp'),
    ('ch02', 'ON TAP CHUONG 2 - DANH GIA CO HOI, Y TUONG PHAT TRIEN KE HOACH KINH DOANH', 2,
     'Đánh giá cơ hội, ý tưởng phát triển kế hoạch kinh doanh'),
    ('ch03', 'ON TAP CHUONG 3 - HUY DONG VON VA CAC CHI SO TAI CHINH COT YEU', 3,
     'Huy động vốn và các chỉ số tài chính cốt yếu'),
    ('ch04', 'ON TAP CHUONG 4 - LUA CHON MO HINH KINH DOANH', 4,
     'Lựa chọn mô hình kinh doanh'),
    ('ch05', 'ON TAP CHUONG 5 - MARKETING CHO DOANH NGHIEP MOI', 5,
     'Marketing cho doanh nghiệp mới'),
]

CHUONG_TIEP_THEO = {
    'ch01': 'Tiếp theo — Chương 2: Đánh giá cơ hội, ý tưởng phát triển kế hoạch kinh doanh',
    'ch02': 'Tiếp theo — Chương 3: Huy động vốn và các chỉ số tài chính cốt yếu',
    'ch03': 'Tiếp theo — Chương 4: Lựa chọn mô hình kinh doanh',
    'ch04': 'Tiếp theo — Chương 5: Marketing cho doanh nghiệp mới',
    'ch05': '',  # chương cuối, không có "tiếp theo"
}


def uoc_tinh_thoi_luong(loi_doc):
    so_tu = len(loi_doc.split())
    return max(THOI_LUONG_TOI_THIEU, so_tu / TOC_DO_DOC * 60.0 + NGHI_CHUYEN_SLIDE)


def anh_xa_slide(cac_diem):
    """Ánh xạ 1-1: điểm đọc thứ i -> ảnh slide số i+1. Mỗi điểm đọc trong
    kich_ban_on_tap.json ứng với đúng một slide trong slide_noi_dung.json,
    cùng thứ tự — xem docstring đầu tệp."""
    return [(i + 1, uoc_tinh_thoi_luong(d['loi_doc']), d['loi_doc'])
            for i, d in enumerate(cac_diem)]


def dung_anh_slide(trinh_duyet, ma_chuong, so_chuong, ten_chuong, slides):
    """Dựng toàn bộ ảnh slide của một chương bằng Playwright, ghi đè
    assets/slides/ch0N/NNN.jpg."""
    thu_muc = os.path.join(GOC, 'assets', 'slides', ma_chuong)
    os.makedirs(thu_muc, exist_ok=True)
    for ten_tep in os.listdir(thu_muc):
        if ten_tep.endswith('.jpg'):
            os.remove(os.path.join(thu_muc, ten_tep))

    tong_so = len(slides)
    trang = trinh_duyet.new_page(viewport={'width': 1280, 'height': 720}, device_scale_factor=1.5)
    for i, s in enumerate(slides):
        so_trang = i + 1
        if so_trang == 1:
            mo_ta = ' '.join(s.get('bullets', []))
            noi_dung = mau.trang_bia(so_chuong, ten_chuong, mo_ta, tong_so)
        elif so_trang == tong_so:
            mo_ta = ' '.join(s.get('bullets', []))
            noi_dung = mau.trang_ket(so_chuong, mo_ta, CHUONG_TIEP_THEO[ma_chuong], so_trang, tong_so)
        else:
            noi_dung = mau.trang_noi_dung(s.get('kicker', ''), s['title'], s['bullets'], so_trang, tong_so)

        trang.set_content(mau.trang_html_day_du(noi_dung))
        trang.wait_for_timeout(80)
        trang.screenshot(path=os.path.join(thu_muc, '%03d.jpg' % so_trang), type='jpeg', quality=90)
    trang.close()


def dung_video(ma_chuong, anh_xa):
    thu_muc_anh = os.path.join(GOC, 'assets', 'slides', ma_chuong)
    duong_dan_concat = '/tmp/%s_concat.txt' % ma_chuong
    dong = []
    for so_anh, dur, _ in anh_xa:
        anh = os.path.join(thu_muc_anh, '%03d.jpg' % so_anh)
        assert os.path.exists(anh), anh
        dong.append("file '%s'" % anh)
        dong.append('duration %.2f' % dur)
    anh_cuoi = os.path.join(thu_muc_anh, '%03d.jpg' % anh_xa[-1][0])
    dong.append("file '%s'" % anh_cuoi)
    with open(duong_dan_concat, 'w', encoding='utf-8') as f:
        f.write('\n'.join(dong))

    duong_dan_ra = os.path.join(
        THU_MUC_VIDEO, 'CHUONG_%s__ON_TAP.mp4' % ma_chuong[2:].lstrip('0'))
    lenh = [
        'ffmpeg', '-y', '-f', 'concat', '-safe', '0', '-i', duong_dan_concat,
        '-vf', 'scale=1280:720,format=yuv420p',
        '-r', '25', '-c:v', 'libx264', '-pix_fmt', 'yuv420p',
        duong_dan_ra,
    ]
    subprocess.run(lenh, check=True, capture_output=True)
    return duong_dan_ra, sum(dur for _, dur, _ in anh_xa)


def dinh_dang_thoi_luong(giay):
    phut, giay_le = divmod(round(giay), 60)
    return '%d phút %02d' % (phut, giay_le) if phut else '%d giây' % giay_le


def them_doan(doc_word, chu, dam=False, nghieng=False, co_chu=11,
              cach_truoc=0, cach_sau=6):
    p = doc_word.add_paragraph()
    p.paragraph_format.space_before = Pt(cach_truoc)
    p.paragraph_format.space_after = Pt(cach_sau)
    r = p.add_run(chu)
    r.bold = dam
    r.italic = nghieng
    r.font.size = Pt(co_chu)
    return p


def dung_pdf_slide():
    """Ghép toàn bộ ảnh slide thật (5 chương, đúng thứ tự) thành một PDF —
    bản gốc để Huong dựng lại ảnh minh họa phong cách 3D đất sét, không phải
    tệp assets/slides/ dùng cho trình xem trong ứng dụng (ảnh gốc giữ nguyên)."""
    anh_list = []
    for ma_chuong, _khoa, _so, _ten in CAC_CHUONG:
        thu_muc = os.path.join(GOC, 'assets', 'slides', ma_chuong)
        so_anh = len([f for f in os.listdir(thu_muc) if f.endswith('.jpg')])
        for i in range(1, so_anh + 1):
            anh_list.append(Image.open(os.path.join(thu_muc, '%03d.jpg' % i)).convert('RGB'))
    duong_dan_ra = os.path.join(THU_MUC_SLIDES, 'ON_TAP_5_CHUONG_EC1606.pdf')
    anh_list[0].save(duong_dan_ra, save_all=True, append_images=anh_list[1:])
    return duong_dan_ra, len(anh_list)


def dung_pdf_kich_ban(cac_dong_tong_hop, noi_dung_chuong):
    """Bản PDF của kịch bản, dựng thẳng bằng reportlab (không qua soffice —
    xem ghi chú đầu tệp), cùng nội dung với bản .docx."""
    pdfmetrics.registerFont(TTFont('VN', PHONG_THUONG))
    pdfmetrics.registerFont(TTFont('VN-Dam', PHONG_DAM))

    duong_dan_ra = os.path.join(THU_MUC_KICH_BAN, 'KICH_BAN_THUYET_MINH__ON_TAP_EC1606.pdf')
    trang_rong, trang_cao = A4
    le = 20 * mm
    rong_chu = trang_rong - 2 * le
    c = canvas.Canvas(duong_dan_ra, pagesize=A4)
    y = [trang_cao - le]

    def viet(chu, dam=False, co_chu=11, cach_sau=4):
        c.setFont('VN-Dam' if dam else 'VN', co_chu)
        for dong in boc_dong(chu, 'VN-Dam' if dam else 'VN', co_chu, rong_chu):
            if y[0] < le:
                c.showPage()
                y[0] = trang_cao - le
                c.setFont('VN-Dam' if dam else 'VN', co_chu)
            c.drawString(le, y[0], dong)
            y[0] -= co_chu + 3
        y[0] -= cach_sau

    def boc_dong(chu, ten_phong, co_chu, rong_toi_da):
        if not chu:
            return ['']
        tu = chu.split(' ')
        dong, dong_hien_tai = [], tu[0]
        for t in tu[1:]:
            thu = dong_hien_tai + ' ' + t
            if pdfmetrics.stringWidth(thu, ten_phong, co_chu) <= rong_toi_da:
                dong_hien_tai = thu
            else:
                dong.append(dong_hien_tai)
                dong_hien_tai = t
        dong.append(dong_hien_tai)
        return dong

    viet('KỊCH BẢN THUYẾT MINH VIDEO ÔN TẬP', dam=True, co_chu=16, cach_sau=4)
    viet('NĂM VIDEO ÔN TẬP — KHỞI SỰ DOANH NGHIỆP (EC1606)', dam=True, co_chu=13, cach_sau=4)
    viet('Học phần: Khởi sự doanh nghiệp (EC1606) — đề cương chi tiết 2025', co_chu=10)
    viet('Giảng viên thuyết minh: Đỗ Thùy Hương – Khoa Kinh doanh và Luật, '
         'Trường Đại học Công nghệ Kỹ thuật Vĩnh Long', co_chu=10)
    viet('Kịch bản cho năm video ôn tập, mỗi video tương ứng một chương, '
         'mỗi video 12-17 phút, dựng theo đúng cấu trúc kịch bản '
         'ec6035tx-chuong-5 của ComDraft (cùng tác giả).', co_chu=10, cach_sau=10)

    viet('HƯỚNG DẪN SỬ DỤNG KỊCH BẢN', dam=True, co_chu=13, cach_sau=4)
    viet('• Nội dung lời đọc lấy nguyên văn từ docs/kich_ban_on_tap.json — '
         'không thêm nội dung mới ngoài đề cương chi tiết và ngân hàng 300 câu hỏi, '
         'trừ các tình huống minh họa được ghi rõ là minh họa (không phải số liệu thật).')
    viet('• Nghỉ ngắn khoảng nửa giây sau mỗi ý; nghỉ 1 giây trước khi chuyển slide.')
    viet('• Thời lượng ước tính ở mỗi slide tính theo tốc độ %d từ/phút — con số này ĐO '
         'THẬT từ bản thu của chính tác giả ở ec6035tx-chuong-5 (ComDraft), KHÔNG PHẢI đo '
         'trên giọng đọc thật của bộ kịch bản này (EnQuiz chưa có bản thu thật). Dùng để '
         'tham khảo khi quay, cần đối chiếu lại tốc độ đọc thật sau khi có bản thu đầu tiên.'
         % TOC_DO_DOC)
    viet('• Mỗi video dài 12-17 phút, đạt chuẩn thời lượng video tự sản xuất theo Phụ lục '
         'Mẫu HLS-VLU-08 mục 4, Quy định đào tạo trực tuyến và xây dựng học liệu số '
         '(QĐ 55/QĐ-ĐHSPKTVL-ĐT_ĐBCL, 07/7/2025).')
    viet('• Năm video câm (chỉ ảnh slide, theo đúng thời lượng ước tính ở trên) đã dựng '
         'sẵn trong videos/ — dùng để tập đọc theo nhịp hoặc để ghi giọng chồng lên (voice-over).')
    viet('• Mỗi video từ Video 2 trở đi mở bằng một câu nhắc lại ý chính chương trước, rồi mới '
         'vào nội dung chương mới; mỗi video khép bằng một đoạn hệ thống hóa 5 câu hỏi tự kiểm '
         'tra và lời nhắc làm 60 câu trắc nghiệm của chương đó trong EnQuiz.', cach_sau=10)

    viet('PHÂN CHIA VIDEO', dam=True, co_chu=13, cach_sau=2)
    for so, ten_chuong, so_slide, tong_thoi_luong, _ in cac_dong_tong_hop:
        viet('VIDEO %d — Chương %d: %s (slide 1–%d, ước tính %s)'
             % (so, so, ten_chuong, so_slide, dinh_dang_thoi_luong(tong_thoi_luong)), co_chu=10)
    y[0] -= 8

    for so, ten_chuong, tong_thoi_luong, so_slide, cac_slide in noi_dung_chuong:
        c.showPage()
        y[0] = trang_cao - le
        viet('VIDEO %d — CHƯƠNG %d: %s' % (so, so, ten_chuong.upper()), dam=True, co_chu=14, cach_sau=2)
        viet('(%d slide, ước tính %s)' % (so_slide, dinh_dang_thoi_luong(tong_thoi_luong)), co_chu=10, cach_sau=8)
        for so_anh, dur, loi_doc, tieu_de_anh in cac_slide:
            viet('SLIDE %d – %s   (~%s)' % (so_anh, tieu_de_anh.upper(), dinh_dang_thoi_luong(dur)),
                 dam=True, co_chu=11, cach_sau=2)
            viet(loi_doc, co_chu=11, cach_sau=8)

    c.save()
    return duong_dan_ra


def main():
    os.makedirs(THU_MUC_VIDEO, exist_ok=True)
    os.makedirs(THU_MUC_KICH_BAN, exist_ok=True)
    os.makedirs(THU_MUC_SLIDES, exist_ok=True)
    kich_ban = json.load(open(KICH_BAN_JSON, encoding='utf-8'))
    slide_noi_dung = json.load(open(SLIDE_NOI_DUNG_JSON, encoding='utf-8'))
    noi_dung_chuong = []  # cho dung_pdf_kich_ban(): (so, ten_chuong, tong_thoi_luong, so_slide, cac_slide)

    doc_word = docx.Document()
    kieu = doc_word.styles['Normal']
    kieu.font.name = 'Calibri'
    kieu.font.size = Pt(11)

    them_doan(doc_word, 'KỊCH BẢN THUYẾT MINH VIDEO ÔN TẬP', dam=True, co_chu=16)
    them_doan(doc_word, 'NĂM VIDEO ÔN TẬP — KHỞI SỰ DOANH NGHIỆP (EC1606)', dam=True, co_chu=13)
    them_doan(doc_word, 'Học phần: Khởi sự doanh nghiệp (EC1606) — đề cương chi tiết 2025', nghieng=True)
    them_doan(doc_word,
              'Giảng viên thuyết minh: Đỗ Thùy Hương – Khoa Kinh doanh và Luật, '
              'Trường Đại học Công nghệ Kỹ thuật Vĩnh Long', nghieng=True)
    them_doan(doc_word,
              'Kịch bản cho năm video ôn tập, mỗi video tương ứng một chương, mỗi video '
              '12-17 phút, dựng theo đúng cấu trúc kịch bản ec6035tx-chuong-5 của ComDraft '
              '(cùng tác giả).', nghieng=True)

    them_doan(doc_word, 'HƯỚNG DẪN SỬ DỤNG KỊCH BẢN', dam=True, co_chu=13, cach_truoc=12)
    them_doan(doc_word, '• Nội dung lời đọc lấy nguyên văn từ docs/kich_ban_on_tap.json — '
                         'không thêm nội dung mới ngoài đề cương chi tiết và ngân hàng 300 câu hỏi, '
                         'trừ các tình huống minh họa được ghi rõ là minh họa (không phải số liệu thật).')
    them_doan(doc_word, '• Nghỉ ngắn khoảng nửa giây sau mỗi ý; nghỉ 1 giây trước khi chuyển slide.')
    them_doan(doc_word,
              '• Thời lượng ước tính ở mỗi slide tính theo tốc độ %d từ/phút — con số này ĐO '
              'THẬT từ bản thu của chính tác giả ở ec6035tx-chuong-5 (ComDraft), KHÔNG PHẢI đo '
              'trên giọng đọc thật của bộ kịch bản này (EnQuiz chưa có bản thu thật). Dùng để '
              'tham khảo khi quay, cần đối chiếu lại tốc độ đọc thật sau khi có bản thu đầu tiên.'
              % TOC_DO_DOC)
    them_doan(doc_word, '• Mỗi video dài 12-17 phút, đạt chuẩn thời lượng video tự sản xuất theo '
                         'Phụ lục Mẫu HLS-VLU-08 mục 4, Quy định đào tạo trực tuyến và xây dựng học '
                         'liệu số (QĐ 55/QĐ-ĐHSPKTVL-ĐT_ĐBCL, 07/7/2025).')
    them_doan(doc_word, '• Năm video câm (chỉ ảnh slide, theo đúng thời lượng ước tính ở trên) đã '
                         'dựng sẵn trong videos/ — dùng để tập đọc theo nhịp hoặc để ghi giọng '
                         'chồng lên (voice-over).')
    them_doan(doc_word, '• Mỗi video từ Video 2 trở đi mở bằng một câu nhắc lại ý chính chương '
                         'trước, rồi mới vào nội dung chương mới; mỗi video khép bằng một đoạn hệ '
                         'thống hóa 5 câu hỏi tự kiểm tra và lời nhắc làm 60 câu trắc nghiệm của '
                         'chương đó trong EnQuiz.')

    them_doan(doc_word, 'PHÂN CHIA VIDEO', dam=True, co_chu=13, cach_truoc=12)
    them_doan(doc_word, '')

    bang = doc_word.add_table(rows=1, cols=4)
    bang.style = 'Light Grid Accent 1'
    tieu_de = bang.rows[0].cells
    for i, h in enumerate(['Video', 'Chương', 'Slide', 'Thời lượng ước tính']):
        tieu_de[i].text = h
        tieu_de[i].paragraphs[0].runs[0].bold = True

    with sync_playwright() as p:
        trinh_duyet = p.chromium.launch(executable_path=CHROME_PATH)

        cac_dong_tong_hop = []
        for ma_chuong, khoa, so, ten_chuong in CAC_CHUONG:
            cac_diem = kich_ban[khoa]
            slides = slide_noi_dung[ma_chuong]
            assert len(cac_diem) == len(slides), (
                '%s: %d điểm đọc nhưng %d slide — hai tệp phải khớp 1-1'
                % (ma_chuong, len(cac_diem), len(slides)))

            dung_anh_slide(trinh_duyet, ma_chuong, so, ten_chuong, slides)

            anh_xa = anh_xa_slide(cac_diem)
            duong_dan_video, tong_thoi_luong = dung_video(ma_chuong, anh_xa)
            cac_dong_tong_hop.append((so, ten_chuong, len(anh_xa), tong_thoi_luong,
                                       os.path.basename(duong_dan_video)))

            doc_word.add_page_break()
            them_doan(doc_word, 'VIDEO %d — CHƯƠNG %d: %s' % (so, so, ten_chuong.upper()),
                      dam=True, co_chu=14)
            them_doan(doc_word, '(%d slide, ước tính %s)' % (len(anh_xa), dinh_dang_thoi_luong(tong_thoi_luong)),
                      nghieng=True, cach_sau=10)

            cac_slide_pdf = []
            for (so_anh, dur, loi_doc), s in zip(anh_xa, slides):
                them_doan(doc_word, 'SLIDE %d – %s   (~%s)' % (so_anh, s['title'].upper(), dinh_dang_thoi_luong(dur)),
                          dam=True, cach_truoc=8)
                them_doan(doc_word, loi_doc)
                cac_slide_pdf.append((so_anh, dur, loi_doc, s['title']))
            noi_dung_chuong.append((so, ten_chuong, tong_thoi_luong, len(anh_xa), cac_slide_pdf))

        trinh_duyet.close()

    doc_word.add_page_break()
    them_doan(doc_word, 'TỔNG HỢP', dam=True, co_chu=13)
    for so, ten_chuong, so_slide, tong_thoi_luong, ten_tep in cac_dong_tong_hop:
        dong = bang.add_row().cells
        dong[0].text = 'VIDEO %d' % so
        dong[1].text = 'Chương %d: %s' % (so, ten_chuong)
        dong[2].text = '1 – %d' % so_slide
        dong[3].text = '~%s' % dinh_dang_thoi_luong(tong_thoi_luong)

    duong_dan_docx = os.path.join(THU_MUC_KICH_BAN, 'KICH_BAN_THUYET_MINH__ON_TAP_EC1606.docx')
    doc_word.save(duong_dan_docx)
    duong_dan_pdf_kich_ban = dung_pdf_kich_ban(cac_dong_tong_hop, noi_dung_chuong)
    duong_dan_pdf_slide, so_anh_slide = dung_pdf_slide()

    print('Đã lưu:', duong_dan_docx)
    print('Đã lưu:', duong_dan_pdf_kich_ban)
    print('Đã lưu:', duong_dan_pdf_slide, '(%d ảnh)' % so_anh_slide)
    for so, ten_chuong, so_slide, tong_thoi_luong, ten_tep in cac_dong_tong_hop:
        print('  Video %d: %s — %d slide, ~%s' % (so, ten_tep, so_slide, dinh_dang_thoi_luong(tong_thoi_luong)))


if __name__ == '__main__':
    main()
